"""Exports Excel (openpyxl) et reçus PDF (WeasyPrint avec fallback HTML)."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.repositories.eleves import EleveRepository, InscriptionRepository
from app.repositories.referentiels import ReferentielRepository
from app.services import dashboard as dash_svc
from app.services import situation as sit_svc

# ------------------------------------------------------------------
# Style Excel
# ------------------------------------------------------------------
_HEADER_FILL = PatternFill("solid", fgColor="0D6EFD")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_TITLE_FONT = Font(bold=True, size=14, color="0D6EFD")
_MONEY_FMT = '#,##0 "Ar"'
_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)


def _apply_header(ws, row: int, headers: list[str]) -> None:
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill = _HEADER_FILL
        c.font = _HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = _BORDER


def _autosize(ws, min_w: int = 10, max_w: int = 40) -> None:
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        length = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
        ws.column_dimensions[letter].width = min(max(length + 2, min_w), max_w)


def _save(wb: Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------
# Export : Dashboard annuel
# ------------------------------------------------------------------
async def export_dashboard_annuel_excel(
    session: AsyncSession, *, annee_id: int | None
) -> bytes:
    if annee_id is None:
        annee = await ReferentielRepository(session).get_annee_active()
        if annee is None:
            raise ValueError("Aucune année scolaire active")
        annee_id = annee.id
    data = await dash_svc.dashboard_annuel(session, annee_id=annee_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Année"
    ws["A1"] = f"SMART KIDS — Bilan annuel {data['libelle']}"
    ws["A1"].font = _TITLE_FONT
    ws.merge_cells("A1:C1")

    row = 3
    for label, val in [
        ("Effectif actif", data["effectif"]),
        ("CA écolage", data["ca_ecolage"]),
        ("CA inscription", data["ca_inscription"]),
        ("CA frais généraux", data["ca_frais_generaux"]),
        ("CA autre", data["ca_autre"]),
        ("CA total brut", data["ca_total_brut"]),
        ("CA retenu (part Jessy)", data["ca_retenu"]),
        ("Versé aux associés", data["verse_associes"]),
        ("Retards à recouvrer", data["retards_a_recouvrer"]),
        ("Avances perçues", data["avances_percues"]),
        ("FG restants dus", data["fg_restants_dus"]),
        ("Droits d'inscription restants dus", data["droits_inscription_restants_dus"]),
        ("Trésorerie actuelle", data["tresorerie_actuelle"]),
        ("Somme manquante cumulée", data["somme_manquante_cumulee"]),
    ]:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        c = ws.cell(row=row, column=2, value=val)
        if isinstance(val, int):
            c.number_format = _MONEY_FMT
        row += 1

    # Tableau mensuel
    ws2 = wb.create_sheet("Mois")
    ws2["A1"] = "Détail mensuel"
    ws2["A1"].font = _TITLE_FONT
    headers = ["Mois", "CA retenu", "Charges", "Salaires",
               "Prélèvements", "Résultat", "Trésorerie fin"]
    _apply_header(ws2, 3, headers)
    r = 4
    for m in data["mois"]:
        ws2.cell(row=r, column=1, value=m["mois_libelle"])
        ws2.cell(row=r, column=2, value=m["ca_retenu"]).number_format = _MONEY_FMT
        ws2.cell(row=r, column=3, value=m["charges"]).number_format = _MONEY_FMT
        ws2.cell(row=r, column=4, value=m["salaires"]).number_format = _MONEY_FMT
        ws2.cell(row=r, column=5, value=m["prelevements"]).number_format = _MONEY_FMT
        ws2.cell(row=r, column=6, value=m["resultat_apres_prelevements"]).number_format = _MONEY_FMT
        ws2.cell(row=r, column=7, value=m["tresorerie_fin"]).number_format = _MONEY_FMT
        r += 1
    _autosize(ws2)
    _autosize(ws)
    return _save(wb)


# ------------------------------------------------------------------
# Export : Situation d'un élève
# ------------------------------------------------------------------
async def export_situation_eleve_excel(
    session: AsyncSession, *, eleve_id: int, annee_id: int | None
) -> bytes:
    eleve = await EleveRepository(session).by_id(eleve_id)
    if eleve is None:
        raise ValueError("Élève introuvable")
    if annee_id is None:
        annee = await ReferentielRepository(session).get_annee_active()
        if annee is None:
            raise ValueError("Aucune année scolaire active")
        annee_id = annee.id

    ins = await InscriptionRepository(session).by_eleve_annee(eleve_id, annee_id)
    if ins is None:
        raise ValueError("Aucune inscription pour cet élève")

    data = await sit_svc.situation_inscription(session, ins)

    wb = Workbook()
    ws = wb.active
    ws.title = "Situation"
    ws["A1"] = f"Situation écolage — {eleve.matricule} {eleve.nom} {eleve.prenoms}"
    ws["A1"].font = _TITLE_FONT
    ws.merge_cells("A1:F1")

    ws["A3"] = "Niveau (id)";       ws["B3"] = data["niveau_id"]
    ws["A4"] = "Écolage mensuel";   ws["B4"] = data["ecolage_mensuel"]; ws["B4"].number_format = _MONEY_FMT
    ws["A5"] = "Remise (%)";        ws["B5"] = data["remise_pct"]
    ws["A6"] = "Droits d'inscription (du)";   ws["B6"] = data["droits_inscription"]["du"]; ws["B6"].number_format = _MONEY_FMT
    ws["A7"] = "Droits d'inscription (payé)"; ws["B7"] = data["droits_inscription"]["paye"]; ws["B7"].number_format = _MONEY_FMT
    ws["A8"] = "Frais généraux (du)";         ws["B8"] = data["frais_generaux"]["du"]; ws["B8"].number_format = _MONEY_FMT
    ws["A9"] = "Frais généraux (payé)";       ws["B9"] = data["frais_generaux"]["paye"]; ws["B9"].number_format = _MONEY_FMT
    ws["A10"] = "Taux recouvrement (%)";      ws["B10"] = data["totaux"]["taux_recouvrement"]

    _apply_header(ws, 12, ["Mois", "Du", "Payé", "Reste dû", "Statut", "Retard"])
    r = 13
    for m in data["mois"]:
        ws.cell(row=r, column=1, value=m["libelle"])
        ws.cell(row=r, column=2, value=m["du"]).number_format = _MONEY_FMT
        ws.cell(row=r, column=3, value=m["paye"]).number_format = _MONEY_FMT
        ws.cell(row=r, column=4, value=m["reste_du"]).number_format = _MONEY_FMT
        ws.cell(row=r, column=5, value=m["statut"])
        ws.cell(row=r, column=6, value="OUI" if m["en_retard"] else "")
        r += 1
    _autosize(ws)
    return _save(wb)


# ------------------------------------------------------------------
# Export : journal des paiements
# ------------------------------------------------------------------
async def export_paiements_excel(
    session: AsyncSession, *, annee_id: int | None,
    date_from: str | None = None, date_to: str | None = None,
) -> bytes:
    if annee_id is None:
        annee = await ReferentielRepository(session).get_annee_active()
        if annee is None:
            raise ValueError("Aucune année scolaire active")
        annee_id = annee.id

    # Cast explicite des bornes pour lever l'ambiguïté de typage d'asyncpg :
    # sans CAST côté gauche, `$2 IS NULL` ne donne pas le type de $2.
    rows = (await session.execute(text("""
        SELECT p.id, p.date_paiement, e.matricule, e.nom, e.prenoms,
               tp.code AS type_code, m.code AS mois_code,
               p.montant, mp.code AS mode_code, p.reference
        FROM paiements p
        JOIN inscriptions i ON i.id = p.inscription_id
        JOIN eleves e ON e.id = i.eleve_id
        JOIN types_paiement tp ON tp.id = p.type_paiement_id
        LEFT JOIN modes_paiement mp ON mp.id = p.mode_paiement_id
        LEFT JOIN mois_scolaires m ON m.id = p.mois_concerne_id
        WHERE p.annee_scolaire_id = :aid
          AND p.parent_paiement_id IS NULL
          AND NOT EXISTS (SELECT 1 FROM paiements a WHERE a.parent_paiement_id = p.id)
          AND (CAST(:from_d AS DATE) IS NULL OR p.date_paiement >= CAST(:from_d AS DATE))
          AND (CAST(:to_d   AS DATE) IS NULL OR p.date_paiement <= CAST(:to_d   AS DATE))
        ORDER BY p.date_paiement, p.id
    """), {"aid": annee_id, "from_d": date_from, "to_d": date_to})).mappings().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Paiements"
    ws["A1"] = "Journal des paiements"
    ws["A1"].font = _TITLE_FONT
    _apply_header(ws, 3, [
        "Id", "Date", "Matricule", "Nom", "Prénoms", "Type",
        "Mois concerné", "Montant", "Mode", "Référence",
    ])
    r = 4
    for row in rows:
        ws.cell(row=r, column=1, value=row["id"])
        ws.cell(row=r, column=2, value=row["date_paiement"])
        ws.cell(row=r, column=3, value=row["matricule"])
        ws.cell(row=r, column=4, value=row["nom"])
        ws.cell(row=r, column=5, value=row["prenoms"])
        ws.cell(row=r, column=6, value=row["type_code"])
        ws.cell(row=r, column=7, value=row["mois_code"])
        ws.cell(row=r, column=8, value=int(row["montant"])).number_format = _MONEY_FMT
        ws.cell(row=r, column=9, value=row["mode_code"])
        ws.cell(row=r, column=10, value=row["reference"])
        r += 1
    ws.cell(row=r + 1, column=7, value="TOTAL").font = Font(bold=True)
    total_cell = ws.cell(row=r + 1, column=8,
                         value=sum(int(x["montant"]) for x in rows))
    total_cell.number_format = _MONEY_FMT
    total_cell.font = Font(bold=True)
    _autosize(ws)
    return _save(wb)

# ------------------------------------------------------------------
# Reçu (HTML + PDF)
# ------------------------------------------------------------------
def _fmt_ar(n: int) -> str:
    return f"{int(n):,}".replace(",", " ") + " Ar"


async def receipt_payload(
    session: AsyncSession, *, paiement_id: int
) -> dict[str, Any]:
    row = (await session.execute(text("""
        SELECT p.id, p.date_paiement, p.montant, p.reference, p.observation,
               tp.code AS type_code, m.code AS mois_code,
               mp.code AS mode_code,
               e.matricule, e.nom, e.prenoms
        FROM paiements p
        LEFT JOIN inscriptions i ON i.id = p.inscription_id
        LEFT JOIN eleves e ON e.id = i.eleve_id
        JOIN types_paiement tp ON tp.id = p.type_paiement_id
        LEFT JOIN modes_paiement mp ON mp.id = p.mode_paiement_id
        LEFT JOIN mois_scolaires m ON m.id = p.mois_concerne_id
        WHERE p.id = :pid
    """), {"pid": paiement_id})).mappings().one_or_none()
    if row is None:
        raise ValueError("Paiement introuvable")
    return dict(row)


def render_receipt_html(payload: dict[str, Any]) -> str:
    date_p = payload["date_paiement"]
    date_str = date_p.strftime("%d/%m/%Y") if date_p else "Non daté"
    montant = _fmt_ar(int(payload["montant"]))
    return f"""<!doctype html>
<html lang="fr"><head><meta charset="utf-8">
<title>Reçu #{payload['id']}</title>
<style>
@page {{ size: A5; margin: 1cm; }}
body {{ font-family: DejaVu Sans, Arial, sans-serif; font-size: 11pt; color: #111; }}
.header {{ text-align: center; margin-bottom: 1cm; }}
.header h1 {{ color: #0d6efd; margin: 0; font-size: 18pt; }}
.header p {{ margin: 2px 0; color: #555; }}
.title {{ font-size: 14pt; font-weight: bold; margin: 0.5cm 0; border-bottom: 2px solid #0d6efd; padding-bottom: 4px; }}
table {{ width: 100%; border-collapse: collapse; margin-top: 0.5cm; }}
td {{ padding: 6px 4px; border-bottom: 1px solid #e5e7eb; }}
td.label {{ color: #666; width: 40%; }}
td.value {{ font-weight: bold; }}
.montant {{ font-size: 16pt; color: #16a34a; font-weight: bold; margin-top: 0.5cm; text-align: right; }}
.footer {{ margin-top: 1.5cm; font-size: 9pt; color: #666; text-align: center; }}
</style></head><body>
<div class="header">
  <h1>SMART KIDS</h1>
  <p>École — Année scolaire 2026-2027</p>
</div>
<div class="title">REÇU N° {payload['id']:06d}</div>
<table>
  <tr><td class="label">Date du paiement</td><td class="value">{date_str}</td></tr>
  <tr><td class="label">Élève</td><td class="value">{payload.get('matricule') or '—'} {payload.get('nom') or ''} {payload.get('prenoms') or ''}</td></tr>
  <tr><td class="label">Type de paiement</td><td class="value">{payload['type_code']}</td></tr>
  <tr><td class="label">Mois concerné</td><td class="value">{payload.get('mois_code') or '—'}</td></tr>
  <tr><td class="label">Mode de paiement</td><td class="value">{payload.get('mode_code') or '—'}</td></tr>
  <tr><td class="label">Référence</td><td class="value">{payload.get('reference') or '—'}</td></tr>
</table>
<div class="montant">Montant : {montant}</div>
<div class="footer">Reçu généré le {datetime.now().strftime('%d/%m/%Y %H:%M')} — Merci de conserver ce document.</div>
</body></html>
"""


def try_pdf(html: str) -> bytes | None:
    """Convertit HTML → PDF via WeasyPrint. Renvoie None si WeasyPrint indisponible."""
    try:
        from weasyprint import HTML  # import tardif
    except Exception:
        return None
    try:
        return HTML(string=html).write_pdf()
    except Exception:
        return None