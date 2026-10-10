"""Import Excel idempotent des 7 feuilles métier.

Chaque ligne est identifiée par un hash SHA-256 de son contenu normalisé + nom de feuille.
Si le hash existe déjà dans import_lignes, la ligne est ignorée (skip).
Sinon, l'entité correspondante est créée via l'ORM, puis le hash est enregistré.

Règles :
- Les dates manquantes deviennent date_a_completer = true → mois_id = NON_DATE.
- Les dates hors année scolaire basculent sur le mois HORS_ANNEE via les triggers DB.
- Aucune ligne n'est supprimée ni mise à jour : l'import est append-only.
"""

from __future__ import annotations

import hashlib
import json
import logging
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from io import BytesIO
from typing import Any, Callable

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.associes import Associe
from app.models.caisse_fg import CaisseFgMouvement
from app.models.charges import Charge
from app.models.eleves import Eleve, Inscription
from app.models.imports import ImportLigne
from app.models.paiements import Paiement
from app.models.parametres import Parametre
from app.models.personnel import Employe
from app.models.prelevements import Prelevement
from app.models.referentiel import (
    MoisScolaire,
    ModePaiement,
    TypePaiement,
)
from app.repositories.referentiels import ReferentielRepository

logger = logging.getLogger(__name__)


# ------------------------------------------------------------------
# Normalisation
# ------------------------------------------------------------------
def _norm(s: Any) -> str:
    if s is None:
        return ""
    t = str(s).strip().lower()
    for src, dst in (("é", "e"), ("è", "e"), ("ê", "e"), ("à", "a"),
                     ("ô", "o"), ("î", "i"), ("û", "u"), ("ç", "c"),
                     ("ù", "u"), ("ï", "i"), ("ö", "o")):
        t = t.replace(src, dst)
    return " ".join(t.split())


ALIASES: dict[str, list[str]] = {
    "matricule":     ["matricule", "mat", "matric"],
    "nom":           ["nom", "nom eleve", "nom de famille"],
    "prenoms":       ["prenoms", "prenom", "prenoms eleve"],
    "niveau":        ["niveau", "classe"],
    "mois":          ["mois", "mois paye", "mois concerne", "periode"],
    "montant":       ["montant", "somme", "valeur"],
    "date":          ["date", "date paiement", "date operation",
                      "date du", "date reglement", "jour"],
    "mode":          ["mode", "mode paiement", "moyen"],
    "reference":     ["reference", "ref", "n° recu", "numero recu", "recu"],
    "observation":   ["observation", "remarque", "commentaire", "notes"],
    "type":          ["type", "type paiement", "nature"],
    "categorie":     ["categorie", "cat"],
    "libelle":       ["libelle", "designation", "description", "intitule"],
    "poste":         ["poste", "fonction"],
    "salaire":       ["salaire", "salaire mensuel", "salaire brut", "brut"],
    "date_embauche": ["date embauche", "embauche"],
    "sens":          ["sens", "entree/sortie"],
    "destinataire":  ["destinataire", "beneficiaire", "motif beneficiaire"],
    "motif":         ["motif", "motif retrait", "raison"],
    "cle":           ["cle", "parametre", "nom parametre"],
    "valeur":        ["valeur", "valeur parametre"],
}


def _find_col(header: list[str], key: str) -> int | None:
    wanted = {_norm(a) for a in ALIASES[key]}
    for i, h in enumerate(header):
        if _norm(h) in wanted:
            return i
    return None


def _cell(row: tuple, idx: int | None) -> Any:
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _to_int(v: Any) -> int | None:
    if v is None or v == "":
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return int(round(v))
    s = str(v).strip().replace(" ", "").replace("\u00a0", "").replace(",", ".")
    try:
        return int(round(float(s)))
    except ValueError:
        return None


def _to_date(v: Any) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


# ------------------------------------------------------------------
# Résolution mois / mode / type / élèves
# ------------------------------------------------------------------
async def _resolve_mois_by_code(
    session: AsyncSession, annee_id: int, value: Any
) -> MoisScolaire | None:
    if value is None:
        return None
    code = str(value).strip()
    m = await session.scalar(
        select(MoisScolaire).where(
            MoisScolaire.annee_scolaire_id == annee_id,
            MoisScolaire.code == code,
        )
    )
    if m is not None:
        return m
    label = _norm(value)
    rows = (await session.scalars(
        select(MoisScolaire).where(
            MoisScolaire.annee_scolaire_id == annee_id,
            MoisScolaire.est_scolaire.is_(True),
        )
    )).all()
    for r in rows:
        if _norm(r.libelle).startswith(label[:4]):
            return r
    return None


async def _resolve_mode(session: AsyncSession, value: Any) -> int | None:
    if value is None:
        return None
    code = _norm(value).upper()
    mapping = {
        "especes": "ESPECES", "cash": "ESPECES",
        "mvola": "MVOLA", "mobile money": "MVOLA",
        "cheque": "CHEQUE",
        "virement": "VIREMENT", "transfert": "VIREMENT",
        "non precise": "NON_PRECISE", "": "NON_PRECISE",
    }
    target = mapping.get(
        code,
        code if code in {"ESPECES", "MVOLA", "CHEQUE", "VIREMENT", "NON_PRECISE"} else None,
    )
    if target is None:
        return None
    m = await session.scalar(
        select(ModePaiement).where(ModePaiement.code == target)
    )
    return m.id if m else None


async def _resolve_type_paiement(session: AsyncSession, value: Any) -> int | None:
    if value is None:
        return None
    code = _norm(value).upper()
    mapping = {
        "ecolage": "ECOLAGE", "scolarite": "ECOLAGE",
        "inscription": "INSCRIPTION", "droit": "INSCRIPTION",
        "droit inscription": "INSCRIPTION",
        "frais generaux": "FRAIS_GENERAUX", "fg": "FRAIS_GENERAUX",
        "autre": "AUTRE", "autres": "AUTRE",
    }
    target = mapping.get(
        code,
        code if code in {"ECOLAGE", "INSCRIPTION", "FRAIS_GENERAUX", "AUTRE"} else None,
    )
    if target is None:
        return None
    t = await session.scalar(select(TypePaiement).where(TypePaiement.code == target))
    return t.id if t else None


async def _resolve_eleve(session: AsyncSession, matricule: Any) -> Eleve | None:
    if not matricule:
        return None
    return await session.scalar(
        select(Eleve).where(Eleve.matricule == str(matricule).strip())
    )


async def _resolve_inscription(
    session: AsyncSession, eleve_id: int, annee_id: int
) -> Inscription | None:
    return await session.scalar(
        select(Inscription).where(
            Inscription.eleve_id == eleve_id,
            Inscription.annee_scolaire_id == annee_id,
        )
    )


async def _resolve_associe_default(session: AsyncSession) -> Associe | None:
    return await session.scalar(
        select(Associe).where(Associe.is_active.is_(True))
        .order_by(Associe.id).limit(1)
    )


# ------------------------------------------------------------------
# Hash
# ------------------------------------------------------------------
def _hash_row(sheet: str, payload: dict[str, Any]) -> str:
    normalized = {
        k: _norm(v) if not isinstance(v, (int, float, bool, type(None))) else v
        for k, v in sorted(payload.items())
    }
    blob = json.dumps({"sheet": sheet, "row": normalized}, sort_keys=True, default=str)
    return hashlib.sha256(blob.encode("utf-8")).hexdigest()


# ------------------------------------------------------------------
# Résultat
# ------------------------------------------------------------------
@dataclass
class SheetResult:
    sheet: str
    total: int = 0
    created: int = 0
    skipped: int = 0
    errors: int = 0
    warnings: list[str] = field(default_factory=list)


@dataclass
class ImportContext:
    session: AsyncSession
    annee_id: int
    created_by: int
    dry_run: bool
    seen_hashes: set[str]
    report: list[SheetResult]


# ------------------------------------------------------------------
# Handlers par feuille
# ------------------------------------------------------------------
async def _handle_suivi_ecolages(
    ctx: ImportContext, rows: list[tuple], header: list[str]
) -> SheetResult:
    r = SheetResult("SUIVI ECOLAGES")
    i_mat = _find_col(header, "matricule")
    i_mois = _find_col(header, "mois")
    i_mnt = _find_col(header, "montant")
    i_date = _find_col(header, "date")
    i_mode = _find_col(header, "mode")
    i_ref = _find_col(header, "reference")
    i_obs = _find_col(header, "observation")

    t_eco = await _resolve_type_paiement(ctx.session, "ECOLAGE")

    for idx, row in enumerate(rows, start=2):
        r.total += 1
        matricule = _cell(row, i_mat)
        montant = _to_int(_cell(row, i_mnt))
        if montant is None or montant <= 0:
            r.errors += 1
            r.warnings.append(f"L{idx}: montant invalide")
            continue
        eleve = await _resolve_eleve(ctx.session, matricule)
        if eleve is None:
            r.errors += 1
            r.warnings.append(f"L{idx}: élève introuvable ({matricule})")
            continue
        ins = await _resolve_inscription(ctx.session, eleve.id, ctx.annee_id)
        if ins is None:
            r.errors += 1
            r.warnings.append(f"L{idx}: aucune inscription pour {matricule}")
            continue
        mois = await _resolve_mois_by_code(
            ctx.session, ctx.annee_id, _cell(row, i_mois)
        )
        if mois is None:
            r.errors += 1
            r.warnings.append(f"L{idx}: mois introuvable ({_cell(row, i_mois)})")
            continue
        if not mois.est_scolaire:
            r.errors += 1
            r.warnings.append(f"L{idx}: mois non scolaire pour un écolage")
            continue

        # Le hash doit contenir tous les champs significatifs.
        # Oublier un champ fait skipper à tort les lignes ayant les autres identiques.
        payload = {
            "matricule": matricule,
            "mois": mois.code,
            "montant": montant,
            "date": _cell(row, i_date),
            "reference": _cell(row, i_ref),
        }
        h = _hash_row(r.sheet, payload)
        if h in ctx.seen_hashes:
            r.skipped += 1
            continue

        dt = _to_date(_cell(row, i_date))
        paiement = Paiement(
            annee_scolaire_id=ctx.annee_id,
            inscription_id=ins.id,
            type_paiement_id=t_eco,
            mois_concerne_id=mois.id,
            montant=montant,
            date_paiement=dt,
            date_a_completer=dt is None,
            mode_paiement_id=await _resolve_mode(ctx.session, _cell(row, i_mode)),
            reference=_cell(row, i_ref),
            observation=_cell(row, i_obs),
            created_by=ctx.created_by,
        )
        if not ctx.dry_run:
            ctx.session.add(paiement)
            await ctx.session.flush()
            ctx.session.add(ImportLigne(
                hash=h, sheet=r.sheet, row_index=idx,
                entity_type="paiements", entity_id=str(paiement.id),
            ))
        ctx.seen_hashes.add(h)
        r.created += 1
    return r


async def _handle_droits_fg(
    ctx: ImportContext, rows: list[tuple], header: list[str]
) -> SheetResult:
    r = SheetResult("DROITS & FG")
    i_mat = _find_col(header, "matricule")
    i_type = _find_col(header, "type")
    i_mnt = _find_col(header, "montant")
    i_date = _find_col(header, "date")
    i_mode = _find_col(header, "mode")
    i_ref = _find_col(header, "reference")

    for idx, row in enumerate(rows, start=2):
        r.total += 1
        matricule = _cell(row, i_mat)
        montant = _to_int(_cell(row, i_mnt))
        type_code = _norm(_cell(row, i_type)).upper() or "INSCRIPTION"
        if montant is None or montant <= 0:
            r.errors += 1
            r.warnings.append(f"L{idx}: montant invalide")
            continue
        eleve = await _resolve_eleve(ctx.session, matricule)
        if eleve is None:
            r.errors += 1
            r.warnings.append(f"L{idx}: élève introuvable ({matricule})")
            continue
        ins = await _resolve_inscription(ctx.session, eleve.id, ctx.annee_id)
        if ins is None:
            r.errors += 1
            r.warnings.append(f"L{idx}: aucune inscription")
            continue
        t_id = await _resolve_type_paiement(ctx.session, type_code)
        if t_id is None:
            r.errors += 1
            r.warnings.append(f"L{idx}: type inconnu ({type_code})")
            continue

        payload = {
            "matricule": matricule,
            "type": type_code,
            "montant": montant,
            "date": _cell(row, i_date),
            "reference": _cell(row, i_ref),
        }
        h = _hash_row(r.sheet, payload)
        if h in ctx.seen_hashes:
            r.skipped += 1
            continue
        dt = _to_date(_cell(row, i_date))
        p = Paiement(
            annee_scolaire_id=ctx.annee_id,
            inscription_id=ins.id,
            type_paiement_id=t_id,
            mois_concerne_id=None,
            montant=montant,
            date_paiement=dt,
            date_a_completer=dt is None,
            mode_paiement_id=await _resolve_mode(ctx.session, _cell(row, i_mode)),
            reference=_cell(row, i_ref),
            created_by=ctx.created_by,
        )
        if not ctx.dry_run:
            ctx.session.add(p)
            await ctx.session.flush()
            ctx.session.add(ImportLigne(
                hash=h, sheet=r.sheet, row_index=idx,
                entity_type="paiements", entity_id=str(p.id),
            ))
        ctx.seen_hashes.add(h)
        r.created += 1
    return r


async def _handle_personnel(
    ctx: ImportContext, rows: list[tuple], header: list[str]
) -> SheetResult:
    r = SheetResult("PERSONNEL")
    i_nom = _find_col(header, "nom")
    i_pre = _find_col(header, "prenoms")
    i_poste = _find_col(header, "poste")
    i_sal = _find_col(header, "salaire")
    i_dt = _find_col(header, "date_embauche")

    for idx, row in enumerate(rows, start=2):
        r.total += 1
        nom = _cell(row, i_nom)
        prenom = _cell(row, i_pre) or ""
        salaire = _to_int(_cell(row, i_sal))
        if not nom or salaire is None or salaire < 0:
            r.errors += 1
            r.warnings.append(f"L{idx}: nom ou salaire invalide")
            continue
        payload = {
            "nom": nom, "prenom": prenom,
            "poste": _cell(row, i_poste),
            "salaire": salaire, "date": _cell(row, i_dt),
        }
        h = _hash_row(r.sheet, payload)
        if h in ctx.seen_hashes:
            r.skipped += 1
            continue
        e = Employe(
            nom=str(nom).strip(), prenom=str(prenom).strip(),
            poste=_cell(row, i_poste),
            salaire_mensuel=salaire,
            date_embauche=_to_date(_cell(row, i_dt)),
            statut="ACTIF",
        )
        if not ctx.dry_run:
            ctx.session.add(e)
            await ctx.session.flush()
            ctx.session.add(ImportLigne(
                hash=h, sheet=r.sheet, row_index=idx,
                entity_type="employes", entity_id=str(e.id),
            ))
        ctx.seen_hashes.add(h)
        r.created += 1
    return r


async def _handle_charges(
    ctx: ImportContext, rows: list[tuple], header: list[str]
) -> SheetResult:
    r = SheetResult("CHARGES FIXES & VARIABLES")
    i_lib = _find_col(header, "libelle")
    i_cat = _find_col(header, "categorie")
    i_mnt = _find_col(header, "montant")
    i_date = _find_col(header, "date")
    i_mode = _find_col(header, "mode")

    for idx, row in enumerate(rows, start=2):
        r.total += 1
        libelle = _cell(row, i_lib)
        montant = _to_int(_cell(row, i_mnt))
        cat = _norm(_cell(row, i_cat)).upper()
        if cat not in {"FIXE", "VARIABLE", "AUTRE"}:
            cat = "VARIABLE"
        if not libelle or montant is None or montant <= 0:
            r.errors += 1
            r.warnings.append(f"L{idx}: libellé ou montant invalide")
            continue
        payload = {
            "libelle": libelle, "categorie": cat,
            "montant": montant, "date": _cell(row, i_date),
        }
        h = _hash_row(r.sheet, payload)
        if h in ctx.seen_hashes:
            r.skipped += 1
            continue
        dt = _to_date(_cell(row, i_date))
        c = Charge(
            annee_scolaire_id=ctx.annee_id,
            mois_id=None,  # rempli par le trigger derive_mois_operation
            libelle=str(libelle).strip(),
            categorie=cat,
            montant=montant,
            date_operation=dt,
            date_a_completer=dt is None,
            mode_paiement_id=await _resolve_mode(ctx.session, _cell(row, i_mode)),
            created_by=ctx.created_by,
        )
        if not ctx.dry_run:
            ctx.session.add(c)
            await ctx.session.flush()
            ctx.session.add(ImportLigne(
                hash=h, sheet=r.sheet, row_index=idx,
                entity_type="charges", entity_id=str(c.id),
            ))
        ctx.seen_hashes.add(h)
        r.created += 1
    return r


async def _handle_retraits(
    ctx: ImportContext, rows: list[tuple], header: list[str]
) -> SheetResult:
    r = SheetResult("RETRAITS JESSY")
    i_mnt = _find_col(header, "montant")
    i_date = _find_col(header, "date")
    i_dest = _find_col(header, "destinataire")
    i_motif = _find_col(header, "motif")
    i_type = _find_col(header, "type")

    jessy = await _resolve_associe_default(ctx.session)
    if jessy is None:
        r.warnings.append("Aucun associé actif — feuille ignorée")
        return r

    for idx, row in enumerate(rows, start=2):
        r.total += 1
        montant = _to_int(_cell(row, i_mnt))
        if montant is None or montant <= 0:
            r.errors += 1
            r.warnings.append(f"L{idx}: montant invalide")
            continue
        type_v = _norm(_cell(row, i_type)).upper() or "RETRAIT"
        if type_v not in {"RETRAIT", "VERSEMENT_FIXE", "SUPPLEMENT"}:
            type_v = "RETRAIT"

        # Le hash doit contenir le motif : deux retraits au même montant,
        # même date, même destinataire, mais motifs différents sont distincts.
        payload = {
            "montant": montant,
            "date": _cell(row, i_date),
            "destinataire": _cell(row, i_dest),
            "type": type_v,
            "motif": _cell(row, i_motif),
        }
        h = _hash_row(r.sheet, payload)
        if h in ctx.seen_hashes:
            r.skipped += 1
            continue
        dt = _to_date(_cell(row, i_date))
        p = Prelevement(
            annee_scolaire_id=ctx.annee_id,
            associe_id=jessy.id,
            type=type_v,
            montant=montant,
            date_prelevement=dt,
            date_a_completer=dt is None,
            destinataire=_cell(row, i_dest),
            motif=_cell(row, i_motif),
            created_by=ctx.created_by,
        )
        if not ctx.dry_run:
            ctx.session.add(p)
            await ctx.session.flush()
            ctx.session.add(ImportLigne(
                hash=h, sheet=r.sheet, row_index=idx,
                entity_type="prelevements", entity_id=str(p.id),
            ))
        ctx.seen_hashes.add(h)
        r.created += 1
    return r


async def _handle_caisse_fg(
    ctx: ImportContext, rows: list[tuple], header: list[str]
) -> SheetResult:
    r = SheetResult("CAISSE FG")
    i_sens = _find_col(header, "sens")
    i_mnt = _find_col(header, "montant")
    i_lib = _find_col(header, "libelle")
    i_date = _find_col(header, "date")

    for idx, row in enumerate(rows, start=2):
        r.total += 1
        montant = _to_int(_cell(row, i_mnt))
        sens = _norm(_cell(row, i_sens)).upper() or "ENTREE"
        if sens not in {"ENTREE", "SORTIE"}:
            r.errors += 1
            r.warnings.append(f"L{idx}: sens invalide ({sens})")
            continue
        if montant is None or montant <= 0:
            r.errors += 1
            r.warnings.append(f"L{idx}: montant invalide")
            continue
        libelle = _cell(row, i_lib) or "(sans libellé)"
        payload = {
            "sens": sens, "montant": montant,
            "libelle": libelle, "date": _cell(row, i_date),
        }
        h = _hash_row(r.sheet, payload)
        if h in ctx.seen_hashes:
            r.skipped += 1
            continue
        dt = _to_date(_cell(row, i_date))
        m = CaisseFgMouvement(
            annee_scolaire_id=ctx.annee_id,
            mois_id=None,
            sens=sens,
            montant=montant,
            date_operation=dt,
            date_a_completer=dt is None,
            libelle=str(libelle).strip(),
            created_by=ctx.created_by,
        )
        if not ctx.dry_run:
            ctx.session.add(m)
            await ctx.session.flush()
            ctx.session.add(ImportLigne(
                hash=h, sheet=r.sheet, row_index=idx,
                entity_type="caisse_fg_mouvements", entity_id=str(m.id),
            ))
        ctx.seen_hashes.add(h)
        r.created += 1
    return r


async def _handle_parametres(
    ctx: ImportContext, rows: list[tuple], header: list[str]
) -> SheetResult:
    r = SheetResult("PARAMETRES")
    i_cle = _find_col(header, "cle")
    i_val = _find_col(header, "valeur")

    for idx, row in enumerate(rows, start=2):
        r.total += 1
        cle = _cell(row, i_cle)
        valeur = _cell(row, i_val)
        if not cle:
            r.errors += 1
            r.warnings.append(f"L{idx}: clé manquante")
            continue
        payload = {"cle": cle, "valeur": valeur}
        h = _hash_row(r.sheet, payload)
        if h in ctx.seen_hashes:
            r.skipped += 1
            continue
        existing = await ctx.session.scalar(
            select(Parametre).where(
                Parametre.annee_scolaire_id == ctx.annee_id,
                Parametre.cle == str(cle).strip(),
            )
        )
        if existing is not None:
            r.skipped += 1
            continue
        if not ctx.dry_run:
            ctx.session.add(Parametre(
                annee_scolaire_id=ctx.annee_id,
                cle=str(cle).strip(),
                valeur=valeur,
                updated_by=ctx.created_by,
            ))
            await ctx.session.flush()
            ctx.session.add(ImportLigne(
                hash=h, sheet=r.sheet, row_index=idx,
                entity_type="parametres", entity_id=str(cle),
            ))
        ctx.seen_hashes.add(h)
        r.created += 1
    return r


HANDLERS: dict[str, Callable] = {
    "SUIVI ECOLAGES": _handle_suivi_ecolages,
    "DROITS & FG": _handle_droits_fg,
    "PERSONNEL": _handle_personnel,
    "CHARGES FIXES & VARIABLES": _handle_charges,
    "RETRAITS JESSY": _handle_retraits,
    "CAISSE FG": _handle_caisse_fg,
    "PARAMETRES": _handle_parametres,
}


# ------------------------------------------------------------------
# Point d'entrée
# ------------------------------------------------------------------
async def import_workbook(
    session: AsyncSession, *,
    file_bytes: bytes,
    file_name: str,
    created_by: int,
    dry_run: bool = False,
) -> dict:
    started = datetime.now(timezone.utc)

    annee = await ReferentielRepository(session).get_annee_active()
    if annee is None:
        raise ValueError("Aucune année scolaire active")

    existing_hashes = set(
        (await session.scalars(select(ImportLigne.hash))).all()
    )

    wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    ctx = ImportContext(
        session=session, annee_id=annee.id, created_by=created_by,
        dry_run=dry_run, seen_hashes=existing_hashes, report=[],
    )

    for sheet_name in wb.sheetnames:
        key = _norm(sheet_name).upper()
        matched = next((k for k in HANDLERS if _norm(k).upper() == key), None)
        if matched is None:
            matched = next(
                (k for k in HANDLERS if _norm(k).upper().startswith(key[:8])),
                None,
            )
        if matched is None:
            ctx.report.append(SheetResult(sheet_name,
                                          warnings=["Feuille non reconnue, ignorée"]))
            continue
        ws = wb[sheet_name]
        raw_rows = list(ws.iter_rows(values_only=True))
        if not raw_rows:
            ctx.report.append(SheetResult(matched))
            continue
        header = [str(c) if c is not None else "" for c in raw_rows[0]]
        body = [r for r in raw_rows[1:] if any(c not in (None, "") for c in r)]
        result = await HANDLERS[matched](ctx, body, header)
        ctx.report.append(result)

    wb.close()

    if not dry_run:
        await session.commit()

    finished = datetime.now(timezone.utc)

    return {
        "file_name": file_name,
        "dry_run": dry_run,
        "sheets": [
            {"sheet": s.sheet, "total": s.total, "created": s.created,
             "skipped": s.skipped, "errors": s.errors, "warnings": s.warnings}
            for s in ctx.report
        ],
        "total_created": sum(s.created for s in ctx.report),
        "total_skipped": sum(s.skipped for s in ctx.report),
        "total_errors": sum(s.errors for s in ctx.report),
        "started_at": started,
        "finished_at": finished,
    }