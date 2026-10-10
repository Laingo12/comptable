"""Exports Excel et reçus (HTML/PDF)."""

from __future__ import annotations

import io
import uuid

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from app.core.security import hash_password
from app.main import app


@pytest_asyncio.fixture
async def client():
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as ac:
        yield ac


@pytest_asyncio.fixture
async def admin_token(session_owner):
    suffix = uuid.uuid4().hex[:8]
    from app.models.user import User as U
    session_owner.add(U(
        username=f"exp_{suffix}", email=None,
        password_hash=hash_password("exp12345"),
        full_name=f"Exp {suffix}", role="ADMIN", is_active=True,
    ))
    await session_owner.commit()
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as ac:
        r = await ac.post("/api/v1/auth/login",
                          json={"username": f"exp_{suffix}", "password": "exp12345"})
        return r.json()["access_token"]


async def _creer_paiement(client, headers, suffix):
    r = await client.post("/api/v1/familles",
                          json={"nom": f"ExpFam {suffix}"}, headers=headers)
    famille_id = r.json()["id"]
    r2 = await client.post("/api/v1/eleves",
                           json={"nom": f"E{suffix}", "prenoms": "T",
                                 "famille_id": famille_id}, headers=headers)
    eleve_id = r2.json()["id"]
    r3 = await client.get("/api/v1/niveaux", headers=headers)
    ps = next(n for n in r3.json() if n["code"] == "PS")
    r4 = await client.get("/api/v1/mois-scolaires", headers=headers)
    m_sept = next(m for m in r4.json() if m["code"] == "2026-09")
    r5 = await client.post("/api/v1/inscriptions",
                           json={"eleve_id": eleve_id, "niveau_id": ps["id"],
                                 "mois_entree_id": m_sept["id"]}, headers=headers)
    ins_id = r5.json()["id"]
    r6 = await client.get("/api/v1/types-paiement", headers=headers)
    t_eco = next(t for t in r6.json() if t["code"] == "ECOLAGE")["id"]
    r7 = await client.post("/api/v1/paiements", json={
        "inscription_id": ins_id, "type_paiement_id": t_eco,
        "mois_concerne_id": m_sept["id"], "montant": 80_000,
        "date_paiement": "2026-09-05",
        "reference": f"REC-{suffix}",
    }, headers=headers)
    return eleve_id, r7.json()["id"]


@pytest.mark.asyncio
async def test_export_dashboard_annuel_xlsx(client, admin_token):
    headers = {"Authorization": f"Bearer {admin_token}"}
    r = await client.get("/api/v1/exports/dashboard-annuel.xlsx", headers=headers)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    # Lecture du contenu
    wb = load_workbook(io.BytesIO(r.content))
    assert "Année" in wb.sheetnames
    assert "Mois" in wb.sheetnames
    ws = wb["Année"]
    assert ws["A1"].value.startswith("SMART KIDS")


@pytest.mark.asyncio
async def test_export_situation_eleve_xlsx(client, admin_token):
    headers = {"Authorization": f"Bearer {admin_token}"}
    suffix = uuid.uuid4().hex[:6]
    eleve_id, _ = await _creer_paiement(client, headers, suffix)

    r = await client.get(f"/api/v1/exports/eleves/{eleve_id}/situation.xlsx",
                         headers=headers)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Situation"]
    assert "Situation écolage" in ws["A1"].value
    # 10 mois listés à partir de la ligne 13
    assert ws.cell(row=13, column=1).value is not None


@pytest.mark.asyncio
async def test_export_paiements_xlsx(client, admin_token):
    headers = {"Authorization": f"Bearer {admin_token}"}
    suffix = uuid.uuid4().hex[:6]
    await _creer_paiement(client, headers, suffix)

    r = await client.get("/api/v1/exports/paiements.xlsx", headers=headers)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Paiements"]
    assert ws["A1"].value == "Journal des paiements"
    # Au moins une ligne
    assert ws.cell(row=4, column=1).value is not None


@pytest.mark.asyncio
async def test_recu_html(client, admin_token):
    headers = {"Authorization": f"Bearer {admin_token}"}
    suffix = uuid.uuid4().hex[:6]
    _, paiement_id = await _creer_paiement(client, headers, suffix)

    r = await client.get(
        f"/api/v1/exports/paiements/{paiement_id}/recu?format=html",
        headers=headers,
    )
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/html")
    body = r.text
    assert "SMART KIDS" in body
    assert f"REÇU N° {paiement_id:06d}" in body
    assert "80 000 Ar" in body
    assert f"REC-{suffix}" in body


@pytest.mark.asyncio
async def test_recu_pdf_ou_fallback(client, admin_token):
    """En prod (Docker) → PDF. En dev Windows sans GTK → HTML de repli."""
    headers = {"Authorization": f"Bearer {admin_token}"}
    suffix = uuid.uuid4().hex[:6]
    _, paiement_id = await _creer_paiement(client, headers, suffix)

    r = await client.get(
        f"/api/v1/exports/paiements/{paiement_id}/recu",
        headers=headers,
    )
    assert r.status_code == 200
    ct = r.headers["content-type"]
    if ct.startswith("application/pdf"):
        assert r.content.startswith(b"%PDF")
    else:
        # Fallback HTML attendu
        assert r.headers.get("x-receipt-fallback") == "html-weasyprint-unavailable"
        assert "SMART KIDS" in r.text